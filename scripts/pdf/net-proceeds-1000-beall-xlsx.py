# -*- coding: utf-8 -*-
"""Formula-driven net proceeds workbook for 1000 Beall Dr (Eric Norman).
Yellow cells are editable inputs; everything else is a formula.
Run: py scripts/pdf/net-proceeds-1000-beall-xlsx.py
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

OUT = r"C:\Users\kylej\Documents\GitHub\friedman-brain\Friedman Brain\listings\1000-beall-dr\Net-Proceeds-1000-Beall-Dr.xlsx"

wb = Workbook(); ws = wb.active; ws.title = "Net Proceeds"
teal = "0F5C63"; gold = "C9A96A"; cream = "FAF8F5"; yellow = "FFF3B0"
HF = Font(bold=True, color="FFFFFF"); H = PatternFill("solid", fgColor=teal)
Y = PatternFill("solid", fgColor=yellow); C = PatternFill("solid", fgColor=cream)
B = Font(bold=True); money = '#,##0;(#,##0)'; pct = '0.00%'
thin = Border(*(Side(style="thin", color="D9D2C4"),)*4)

ws.column_dimensions["A"].width = 46
for col in ("B","C","D"): ws.column_dimensions[col].width = 16

ws["A1"] = "1000 Beall Dr, Joppa MD 21085 - Estimated Net Proceeds"; ws["A1"].font = Font(bold=True, size=13, color=teal)
ws["A2"] = "The Friedman Team | Kyle Friedman 443-789-3101 | yellow cells are editable"; ws["A2"].font = Font(italic=True, color="5B6B6E", size=9)

r = 4
ws.cell(r,1,"Scenario").fill = H; ws.cell(r,1).font = HF
for i,label in enumerate(("Conservative","Recommended","Speculative")):
    ws.cell(r,2+i, label).fill = H; ws.cell(r,2+i).font = HF

r += 1
ws.cell(r,1,"Sale price").font = B
for i,v in enumerate((300900,329900,364900)):
    cell = ws.cell(r,2+i,v); cell.fill = Y; cell.number_format = money
price_row = r

r += 1
ws.cell(r,1,"Commission %").font = B
for i in range(3):
    cell = ws.cell(r,2+i,0.05); cell.fill = Y; cell.number_format = pct
comm_row = r

r += 1
ws.cell(r,1,"MD state + Harford County transfer & recordation, seller share %").font = B
for i in range(3):
    cell = ws.cell(r,2+i,0.0108); cell.fill = Y; cell.number_format = pct
xfer_row = r

r += 1
ws.cell(r,1,"Settlement / closing fee")
for i in range(3):
    cell = ws.cell(r,2+i,500); cell.fill = Y; cell.number_format = money
fee_row = r

r += 1
ws.cell(r,1,"Deed & doc prep")
for i in range(3):
    cell = ws.cell(r,2+i,200); cell.fill = Y; cell.number_format = money
deed_row = r

r += 1
ws.cell(r,1,"Lien release, both loans")
for i in range(3):
    cell = ws.cell(r,2+i,250); cell.fill = Y; cell.number_format = money
lien_row = r

r += 1
ws.cell(r,1,"Wire, courier & notary")
for i in range(3):
    cell = ws.cell(r,2+i,150); cell.fill = Y; cell.number_format = money
wire_row = r

r += 1
ws.cell(r,1,"Termite / WDI inspection")
for i in range(3):
    cell = ws.cell(r,2+i,100); cell.fill = Y; cell.number_format = money
term_row = r

r += 1
ws.cell(r,1,"Seller concession / repairs / warranty (optional)")
for i in range(3):
    cell = ws.cell(r,2+i,0); cell.fill = Y; cell.number_format = money
conc_row = r

r += 1
ws.cell(r,1,"Total estimated costs").font = B
for i,col in enumerate(("B","C","D")):
    f = (f"={col}{price_row}*{col}{comm_row}+{col}{price_row}*{col}{xfer_row}"
         f"+{col}{fee_row}+{col}{deed_row}+{col}{lien_row}+{col}{wire_row}+{col}{term_row}+{col}{conc_row}")
    cell = ws.cell(r,2+i,f); cell.number_format = money; cell.font = B
tot_row = r

r += 1
ws.cell(r,1,"NET BEFORE LOAN PAYOFFS").font = Font(bold=True, color=teal)
for i,col in enumerate(("B","C","D")):
    cell = ws.cell(r,2+i, f"={col}{price_row}-{col}{tot_row}")
    cell.number_format = money; cell.font = Font(bold=True, color=teal); cell.fill = PatternFill("solid", fgColor=gold)
netbp_row = r

r += 1
ws.cell(r,1,"Net as % of sale price")
for i,col in enumerate(("B","C","D")):
    cell = ws.cell(r,2+i, f"={col}{netbp_row}/{col}{price_row}"); cell.number_format = pct

r += 1
ws.cell(r,1,"Estimated loan balance (not a payoff statement)").font = B
for i in range(3):
    cell = ws.cell(r,2+i,234785); cell.fill = Y; cell.number_format = money
payoff_row = r

r += 1
ws.cell(r,1,"ESTIMATED NET TO YOU").font = Font(bold=True, color=teal, size=12)
for i,col in enumerate(("B","C","D")):
    cell = ws.cell(r,2+i, f"={col}{netbp_row}-{col}{payoff_row}")
    cell.number_format = money; cell.font = Font(bold=True, color=teal, size=12); cell.fill = C

r += 2
for line in [
  "Assumptions:",
  "- Estimate only, not a closing disclosure.",
  "- Commission 5.0% is a placeholder, negotiable, set in the listing agreement.",
  "- Harford County transfer tax 1.0% + MD state transfer 0.5% + Harford recordation ~$6.60/$1,000 (0.66%),",
  "  shown at seller customary half share (1.08% combined). Exact split per contract, confirm with title.",
  "- Estimated loan balance: $234,785 (Sept 2026). This is an estimate, not a payoff statement, actual",
  "  payoff may differ slightly with per-diem interest and fees. Confirm with a current statement before closing.",
  "- List range $300,900-$364,900 reflects the 21085 market-evidence pricing (6 closed comps in Joppa,",
  "  average closed price $397,056, RPR CMA, Sept 2026), adjusted down for the home's dated kitchen and",
  "  bathrooms, offset by the largest lot in the comp set and recently updated roof, HVAC, and water heater.",
  "- Section 121: primary residence 2 of last 5 yrs excludes first 250k gain (single) / 500k (MFJ) from",
  "  federal tax. Purchased 2017 for $175,000, so gain is well under that threshold at every scenario.",
  "  Confirm with CPA.",
]:
    ws.cell(r,1,line).font = Font(italic=True, color="5B6B6E", size=9); r += 1

wb.save(OUT); print("wrote", OUT)
