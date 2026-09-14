# -*- coding: utf-8 -*-
"""Formula-driven net proceeds workbook for 261 Magothy Bridge Rd (Blaine Welker).
Yellow cells are editable inputs; everything else is a formula.
Run: py scripts/pdf/net-proceeds-261-magothy-xlsx.py
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

OUT = r"C:\Users\kylej\Documents\GitHub\friedman-brain\Friedman Brain\listings\261-magothy-bridge-rd\Net-Proceeds-261-Magothy-Bridge.xlsx"

wb = Workbook(); ws = wb.active; ws.title = "Net Proceeds"
teal = "0F5C63"; gold = "C9A96A"; cream = "FAF8F5"; yellow = "FFF3B0"
HF = Font(bold=True, color="FFFFFF"); H = PatternFill("solid", fgColor=teal)
Y = PatternFill("solid", fgColor=yellow); C = PatternFill("solid", fgColor=cream)
B = Font(bold=True); money = '#,##0;(#,##0)'; pct = '0.00%'
thin = Border(*(Side(style="thin", color="D9D2C4"),)*4)

ws.column_dimensions["A"].width = 42
for col in ("B","C","D"): ws.column_dimensions[col].width = 16

ws["A1"] = "261 Magothy Bridge Rd, Pasadena MD 21122 - Estimated Net Proceeds"; ws["A1"].font = Font(bold=True, size=13, color=teal)
ws["A2"] = "The Friedman Team | Kyle Friedman 443-789-3101 | yellow cells are editable"; ws["A2"].font = Font(italic=True, color="5B6B6E", size=9)

r = 4
ws.cell(r,1,"Scenario").fill = H; ws.cell(r,1).font = HF
for i,c in enumerate(("B","C","D")):
    ws.cell(r,2+i, f"List {i+1}").fill = H; ws.cell(r,2+i).font = HF

r += 1
ws.cell(r,1,"Sale price").font = B
for i,v in enumerate((475000,490000,510000)):
    cell = ws.cell(r,2+i,v); cell.fill = Y; cell.number_format = money
price_row = r

r += 1
ws.cell(r,1,"Commission %").font = B
for i in range(3):
    cell = ws.cell(r,2+i,0.05); cell.fill = Y; cell.number_format = pct
comm_row = r

r += 1
ws.cell(r,1,"AA transfer + recordation, seller share %").font = B
for i in range(3):
    cell = ws.cell(r,2+i,0.010); cell.fill = Y; cell.number_format = pct
xfer_row = r

r += 1
ws.cell(r,1,"Settlement / closing fee")
for i in range(3):
    cell = ws.cell(r,2+i,500); cell.fill = Y; cell.number_format = money
fee_row = r

r += 1
ws.cell(r,1,"Deed & doc prep")
for i in range(3):
    cell = ws.cell(r,2+i,250); cell.fill = Y; cell.number_format = money
deed_row = r

r += 1
ws.cell(r,1,"Payoff processing & lien release")
for i in range(3):
    cell = ws.cell(r,2+i,200); cell.fill = Y; cell.number_format = money
lien_row = r

r += 1
ws.cell(r,1,"Wire, courier & notary")
for i in range(3):
    cell = ws.cell(r,2+i,150); cell.fill = Y; cell.number_format = money
wire_row = r

r += 1
ws.cell(r,1,"Termite inspection")
for i in range(3):
    cell = ws.cell(r,2+i,75); cell.fill = Y; cell.number_format = money
term_row = r

r += 1
ws.cell(r,1,"Seller concession / repairs / warranty (optional)")
for i in range(3):
    cell = ws.cell(r,2+i,0); cell.fill = Y; cell.number_format = money
conc_row = r

r += 1
ws.cell(r,1,"Total estimated costs").font = B
for i,col in enumerate(("B","C","D")):
    f = (f"={col}{price_row}*{col}{comm_row}+{col}{price_row}*{col}{xfer_row}"
         f"+{col}{fee_row}+{col}{deed_row}+{col}{lien_row}+{col}{wire_row}+{col}{term_row}+{col}{conc_row}")
    cell = ws.cell(r,2+i,f); cell.number_format = money; cell.font = B
tot_row = r

r += 1
ws.cell(r,1,"NET BEFORE LOAN PAYOFF").font = Font(bold=True, color=teal)
for i,col in enumerate(("B","C","D")):
    cell = ws.cell(r,2+i, f"={col}{price_row}-{col}{tot_row}")
    cell.number_format = money; cell.font = Font(bold=True, color=teal); cell.fill = PatternFill("solid", fgColor=gold)
netbp_row = r

r += 1
ws.cell(r,1,"Net as % of sale price")
for i,col in enumerate(("B","C","D")):
    cell = ws.cell(r,2+i, f"={col}{netbp_row}/{col}{price_row}"); cell.number_format = pct
r += 1
ws.cell(r,1,"Loan and lien payoff (enter your statement totals)").font = B
for i in range(3):
    cell = ws.cell(r,2+i,0); cell.fill = Y; cell.number_format = money
payoff_row = r

r += 1
ws.cell(r,1,"ESTIMATED NET TO YOU").font = Font(bold=True, color=teal, size=12)
for i,col in enumerate(("B","C","D")):
    cell = ws.cell(r,2+i, f"={col}{netbp_row}-{col}{payoff_row}")
    cell.number_format = money; cell.font = Font(bold=True, color=teal, size=12); cell.fill = C

r += 2
for line in [
  "Assumptions:",
  "- Estimate only, not a closing disclosure.",
  "- Commission 5.0% is a placeholder, negotiable, set in the listing agreement.",
  "- AA County transfer + recordation shown at seller customary ~1.0%; exact split per contract, confirm with title.",
  "- Loan/lien payoff: RPR shows a first mortgage, a HELOC, and a 2016 commercial line recorded over the years. Enter current payoffs for the ones still open.",
  "- List range 475k-510k brackets the competitive zone (21122 median sold 478,500, ~272/sqft, Aug 2026). Exact price set at the walk-through.",
  "- Section 121: primary residence 2 of last 5 yrs excludes first 250k gain (single) / 500k (MFJ) from federal tax. Confirm with CPA.",
]:
    ws.cell(r,1,line).font = Font(italic=True, color="5B6B6E", size=9); r += 1

wb.save(OUT); print("wrote", OUT)
